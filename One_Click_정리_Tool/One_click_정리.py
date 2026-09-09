"""One Click 데이터 정리 도구.

CA410(색좌표) / LMK6(BU·WU 균일도) 측정 원본을 읽어
- Panel_ID 기준으로 최신 측정만 남긴 개별 정리 파일(2개 시트)
- 두 데이터를 병합해 양식 엑셀에 기입한 통합 정리 파일
을 생성한다.
"""

import io
import math
import os
import sys
import threading
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, scrolledtext

import pandas as pd
from openpyxl import load_workbook

APP_VERSION = "1.4.1"
APP_TITLE = f"One Click 데이터 정리 도구 v{APP_VERSION}"

# [설정] 지원 시간 포맷 (앞쪽 포맷부터 순서대로 시도)
SUPPORTED_TIME_FORMATS = (
    "%Y.%m.%d %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
)

# CSV 인코딩 자동 판별 순서 (국내 설비 로그는 cp949/euc-kr 인 경우가 많다)
CSV_ENCODINGS = ("utf-8-sig", "utf-8", "cp949", "euc-kr")

# [설정] 양식 파일 탐색
# 1순위: 실행 파일(또는 스크립트)과 같은 폴더에 동봉된 양식
# 2순위: 아래 상대 경로 후보 (기존 사내 폴더 구조 호환용)
TEMPLATE_FILENAME = "양식_Daily data_BU설비_광학데이터_날짜.xlsx"
LEGACY_TEMPLATE_RELPATHS = (
    os.path.join(
        "..", "..", "업무", "00. 곽원용", "01. PJT", "2025",
        "3. 12.9 inch_VDA129BLAC01", "0. 양산 대응 폴더", "Daily data 정리",
        TEMPLATE_FILENAME,
    ),
)

# [설정] 양식 기입 위치 (1-based). 데이터는 TEMPLATE_START_ROW 행부터 채운다.
TEMPLATE_START_ROW = 20
TEMPLATE_NO_COLUMN = 2
TEMPLATE_PANEL_COLUMN = 3
TEMPLATE_COLUMN_MAP = {
    "Red_X": 6, "Red_Y": 7,
    "Green_X": 8, "Green_Y": 9,
    "Blue_X": 10, "Blue_Y": 11,
    "White_SX": 13, "White_SY": 14, "White_LV": 15,
    "Black_LV": 16,
    "White_Uniformity": 19, "Black_Uniformity": 20,
    "EEPROM_59": 27,
}

# 시트별 선별 컬럼 정의
CA_SELECTED_COLS = [
    "Time", "Panel", "Panel_ID",
    "Red_X", "Red_Y",
    "Green_X", "Green_Y",
    "Blue_X", "Blue_Y",
    "White_SX", "White_SY", "White_LV",
    "Black_LV", "EEPROM_59",
]

LMK_SELECTED_COLS = [
    "Time", "Panel", "Panel_ID", "Judge",
    "Black_Uniformity", "White_Uniformity",
]

# LMK6에서 통합 양식으로 가져올 컬럼 (Panel_ID 는 병합 키)
LMK_MERGE_COLS = ("Panel_ID", "Black_Uniformity", "White_Uniformity")

# 내부 정렬용 임시 컬럼 (결과 파일에는 남기지 않는다)
_TIME_SORT_COL = "__time_sort"


def excel_value(value) -> "str | float | int | bool | datetime | None":
    """openpyxl 셀에 그대로 넣을 수 있는 값으로 변환한다.

    NaN/NaT 은 None(빈 셀)으로, numpy 스칼라는 파이썬 기본형으로 바꾼다.
    변환하지 않으면 결측치가 엑셀에 'nan' 문자열처럼 남는다.
    """
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, float):
        return None if math.isnan(value) else value
    if isinstance(value, (str, int, bool, datetime)):
        return value

    # numpy 스칼라 등은 .item() 으로 파이썬 기본형으로 내린다.
    item = getattr(value, "item", None)
    if callable(item):
        try:
            unwrapped = item()
        except (ValueError, TypeError):
            return str(value)
        if isinstance(unwrapped, float):
            return None if math.isnan(unwrapped) else unwrapped
        if isinstance(unwrapped, (str, int, bool, datetime)):
            return unwrapped
        return str(unwrapped)
    return str(value)


def parse_time_series(series: pd.Series) -> pd.Series:
    """여러 시간 포맷을 순서대로 적용해 datetime 시리즈로 변환한다.

    이미 변환된 행은 다시 파싱하지 않으므로 포맷 개수와 무관하게 1패스에 가깝다.
    """
    parsed = pd.to_datetime(series, format=SUPPORTED_TIME_FORMATS[0], errors="coerce")
    for fmt in SUPPORTED_TIME_FORMATS[1:]:
        missing = parsed.isna()
        if not missing.any():
            break
        parsed.loc[missing] = pd.to_datetime(
            series.loc[missing], format=fmt, errors="coerce"
        )
    return parsed


def normalize_panel_ids(series: pd.Series) -> pd.Series:
    """Panel_ID 표기 흔들림(공백·구분자·엑셀 .0 접미사·대소문자)을 정규화한다.

    체이닝 대신 정규식 2회로 처리해 중간 Series 생성을 줄였다.
    """
    ids = series.astype(str).str.strip()
    ids = ids.str.replace(r"[\s_-]+", "", regex=True)
    # 엑셀이 숫자로 인식해 붙인 소수점 접미사만 제거한다.
    # (기존 구현은 ID 중간의 '.0' 까지 지워 'AB.05' -> 'AB5' 로 망가뜨렸다)
    ids = ids.str.replace(r"\.0$", "", regex=True)
    return ids.str.upper()


class OneClickApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("800x700")

        self.last_saved_file = None
        self.last_saved_folder = None
        self.template_path = self.get_default_template_path()
        self._main_thread = threading.current_thread()

        self.setup_ui()
        self.log("프로그램이 시작되었습니다.")
        self.check_template()

    def setup_ui(self):
        # 파일 선택 프레임
        file_frame = tk.LabelFrame(self.root, text="입력 파일 설정 (CSV/Excel)", padx=10, pady=10)
        file_frame.pack(fill="x", padx=10, pady=5)

        # CA410
        tk.Label(file_frame, text="1. CA410 (색좌표):").grid(row=0, column=0, sticky="w")
        self.ent_ca = tk.Entry(file_frame)
        self.ent_ca.grid(row=0, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="찾기", command=lambda: self.browse_file(self.ent_ca, "CA410")).grid(row=0, column=2)

        # LMK6
        tk.Label(file_frame, text="2. LMK6 (BU/WU):").grid(row=1, column=0, sticky="w", pady=5)
        self.ent_lmk = tk.Entry(file_frame)
        self.ent_lmk.grid(row=1, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="찾기", command=lambda: self.browse_file(self.ent_lmk, "LMK6")).grid(row=1, column=2)

        # 양식 파일
        tk.Label(file_frame, text="3. 양식 파일:").grid(row=2, column=0, sticky="w")
        self.ent_temp = tk.Entry(file_frame)
        self.ent_temp.insert(0, self.template_path or "")
        self.ent_temp.grid(row=2, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="변경", command=self.browse_template).grid(row=2, column=2)

        # 저장 경로
        tk.Label(file_frame, text="4. 저장 폴더:").grid(row=3, column=0, sticky="w", pady=5)
        self.ent_save = tk.Entry(file_frame)
        self.ent_save.grid(row=3, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="찾기", command=self.browse_folder).grid(row=3, column=2)

        # 옵션 설정
        self.var_make_individual = tk.BooleanVar(value=True)
        self.chk_individual = tk.Checkbutton(file_frame, text="개별 정리 파일(Cleaned) 생성 (2개 시트 구성)", variable=self.var_make_individual)
        self.chk_individual.grid(row=4, column=0, columnspan=3, sticky="w", pady=5)

        file_frame.columnconfigure(1, weight=1)

        # 실행 버튼
        self.btn_run = tk.Button(self.root, text="▶ One-Click 정리 시작",
                                 bg="#28a745", fg="white", font=("맑은 고딕", 12, "bold"),
                                 height=2, command=self.start_process)
        self.btn_run.pack(fill="x", padx=10, pady=10)

        # 로그창
        log_frame = tk.LabelFrame(self.root, text="처리 로그", padx=10, pady=10)
        log_frame.pack(fill="both", expand=True, padx=10, pady=5)
        self.txt_log = scrolledtext.ScrolledText(log_frame, height=15, state="disabled", font=("Consolas", 10))
        self.txt_log.pack(fill="both", expand=True)

        # 결과 버튼
        res_frame = tk.Frame(self.root)
        res_frame.pack(fill="x", padx=10, pady=10)
        self.btn_open_f = tk.Button(res_frame, text="결과 파일 열기", state="disabled", command=self.open_last_file)
        self.btn_open_f.pack(side="left", expand=True, fill="x", padx=5)
        self.btn_open_d = tk.Button(res_frame, text="폴더 열기", state="disabled", command=self.open_last_folder)
        self.btn_open_d.pack(side="left", expand=True, fill="x", padx=5)

    def log(self, msg):
        line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n"
        if threading.current_thread() is self._main_thread:
            self._append_log(line)
        else:
            self.root.after(0, self._append_log, line)

    def _append_log(self, line):
        self.txt_log.config(state="normal")
        self.txt_log.insert(tk.END, line)
        self.txt_log.see(tk.END)
        self.txt_log.config(state="disabled")

    def _set_run_button(self, state, text=None):
        self.btn_run.config(state=state, text=text if text is not None else self.btn_run.cget("text"))

    def _set_result_buttons(self, enabled: bool):
        state = "normal" if enabled else "disabled"
        self.btn_open_f.config(state=state)
        self.btn_open_d.config(state=state)

    def _show_info(self, title, text):
        messagebox.showinfo(title, text)

    def _show_error(self, title, text):
        messagebox.showerror(title, text)

    def get_default_template_path(self):
        """양식 파일 기본 경로를 찾는다. 없으면 빈 문자열."""
        base = os.path.dirname(sys.executable) if getattr(sys, "frozen", False) else os.path.dirname(os.path.abspath(__file__))
        search_roots = (base, os.path.dirname(base))

        # 1순위: 프로그램과 동봉된 양식 파일
        for root in search_roots:
            candidate = os.path.join(root, TEMPLATE_FILENAME)
            if os.path.exists(candidate):
                return candidate

        # 2순위: 기존 사내 폴더 구조 상대 경로
        for root in search_roots:
            for relpath in LEGACY_TEMPLATE_RELPATHS:
                candidate = os.path.normpath(os.path.join(root, relpath))
                if os.path.exists(candidate):
                    return candidate

        return ""

    def check_template(self):
        path = self.ent_temp.get()
        if path and os.path.exists(path):
            self.log(f"양식 파일 감지됨: {os.path.basename(path)}")
            if not self.ent_save.get():
                self.ent_save.insert(0, os.path.dirname(path))
        else:
            self.log("⚠ 기본 경로에 양식 파일이 없습니다. 직접 선택해 주세요.")

    def browse_file(self, entry, name):
        p = filedialog.askopenfilename(title=f"{name} 파일 선택", filetypes=[("Data files", "*.csv *.xlsx *.xls")])
        if p:
            entry.delete(0, tk.END)
            entry.insert(0, p)
            if not self.ent_save.get():
                self.ent_save.insert(0, os.path.dirname(p))

    def browse_template(self):
        p = filedialog.askopenfilename(title="양식 파일 선택", filetypes=[("Excel files", "*.xlsx")])
        if p:
            self.ent_temp.delete(0, tk.END)
            self.ent_temp.insert(0, p)

    def browse_folder(self):
        p = filedialog.askdirectory(title="저장 폴더 선택")
        if p:
            self.ent_save.delete(0, tk.END)
            self.ent_save.insert(0, p)

    def read_df(self, path):
        """CSV/Excel 을 DataFrame 으로 읽는다.

        CSV 는 파일을 한 번만 읽어 메모리 버퍼에서 인코딩을 재시도한다.
        (네트워크 드라이브에서 같은 파일을 4번 다시 읽지 않도록)
        """
        ext = os.path.splitext(path)[1].lower()
        if ext != ".csv":
            return pd.read_excel(path)

        with open(path, "rb") as f:
            raw = f.read()

        last_error = None
        for encoding in CSV_ENCODINGS:
            try:
                return pd.read_csv(io.BytesIO(raw), encoding=encoding)
            except UnicodeDecodeError as exc:
                last_error = exc
        raise UnicodeError(
            f"지원하는 인코딩({', '.join(CSV_ENCODINGS)})으로 읽지 못했습니다: {path}"
        ) from last_error

    def clean_df(self, df):
        """시간순 정렬 후 Panel_ID 기준으로 최신 측정만 남긴다."""
        if "Time" in df.columns:
            parsed = parse_time_series(df["Time"])
            df = df.assign(**{_TIME_SORT_COL: parsed})
            df = df[df[_TIME_SORT_COL].notna()].sort_values(_TIME_SORT_COL, kind="stable")

        if "Panel_ID" in df.columns:
            df = df.assign(Panel_ID=normalize_panel_ids(df["Panel_ID"]))
            df = df[df["Panel_ID"] != ""]
            # 시간 오름차순이므로 keep="last" 가 최신 측정이다.
            df = df.drop_duplicates(subset=["Panel_ID"], keep="last")
            df = df.sort_values("Panel_ID", kind="stable")

        # 내부 정렬용 컬럼은 결과에 남기지 않는다.
        if _TIME_SORT_COL in df.columns:
            df = df.drop(columns=[_TIME_SORT_COL])
        return df

    def save_individual_file(self, df, mode, save_dir):
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = os.path.join(save_dir, f"{mode}_cleaned_{timestamp}.xlsx")

        # 선별 컬럼 중 실제 데이터에 존재하는 것만 사용 (에러 방지)
        selected_cols = CA_SELECTED_COLS if mode == "CA410" else LMK_SELECTED_COLS
        actual_cols = [c for c in selected_cols if c in df.columns]

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Sheet1_All", index=False)
            df[actual_cols].to_excel(writer, sheet_name="Sheet2_Selected", index=False)

        return out_path

    def fill_template(self, merged, template_path, save_dir):
        """병합 결과를 양식 엑셀에 기입하고 저장 경로를 돌려준다."""
        wb = load_workbook(template_path)
        ws = wb.active
        if ws is None:
            raise ValueError(f"양식 파일에 활성 시트가 없습니다: {template_path}")

        # 컬럼 위치를 루프 밖에서 한 번만 계산한다.
        # (기존 구현은 행마다 13개 컬럼에 hasattr 를 호출했다)
        columns = list(merged.columns)
        panel_pos = columns.index("Panel_ID")
        targets = [
            (columns.index(name), col)
            for name, col in TEMPLATE_COLUMN_MAP.items()
            if name in columns
        ]

        for offset, values in enumerate(merged.itertuples(index=False, name=None)):
            row_idx = TEMPLATE_START_ROW + offset
            ws.cell(row=row_idx, column=TEMPLATE_NO_COLUMN, value=offset + 1)
            ws.cell(row=row_idx, column=TEMPLATE_PANEL_COLUMN, value=excel_value(values[panel_pos]))
            for pos, col in targets:
                ws.cell(row=row_idx, column=col, value=excel_value(values[pos]))

        out_path = os.path.join(
            save_dir, f"Daily data 통합정리_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        wb.save(out_path)
        return out_path

    def _collect_inputs(self):
        """입력값을 읽고 검증한다. 문제가 있으면 경고 후 None 을 돌려준다.

        Tk 위젯 접근과 messagebox 는 메인 스레드에서만 안전하므로
        검증은 반드시 작업 스레드 시작 전에 수행한다.
        """
        ca_path = self.ent_ca.get().strip()
        lmk_path = self.ent_lmk.get().strip()
        save_dir = self.ent_save.get().strip()
        temp_path = self.ent_temp.get().strip()

        if not ca_path and not lmk_path:
            messagebox.showwarning("알림", "정리할 파일을 하나 이상 선택해 주세요.")
            return None
        if not save_dir:
            messagebox.showwarning("알림", "저장 폴더를 지정해 주세요.")
            return None
        if not os.path.isdir(save_dir):
            messagebox.showwarning("알림", f"저장 폴더가 존재하지 않습니다:\n{save_dir}")
            return None
        if ca_path and lmk_path and not os.path.exists(temp_path):
            messagebox.showwarning("알림", "통합 정리를 하려면 양식 파일이 필요합니다.")
            return None

        return {
            "ca_path": ca_path,
            "lmk_path": lmk_path,
            "save_dir": save_dir,
            "temp_path": temp_path,
            "make_individual": self.var_make_individual.get(),
        }

    def start_process(self):
        params = self._collect_inputs()
        if params is None:
            # 검증 실패 시 버튼이 '처리 중...' 으로 잠기지 않도록 여기서 종료한다.
            return
        self.btn_run.config(state="disabled", text="처리 중...")
        self._set_result_buttons(False)
        threading.Thread(target=self.run, args=(params,), daemon=True).start()

    def run(self, params):
        ca_path = params["ca_path"]
        lmk_path = params["lmk_path"]
        save_dir = params["save_dir"]
        temp_path = params["temp_path"]
        make_indiv = params["make_individual"]
        both = bool(ca_path and lmk_path)

        try:
            df_ca = df_lmk = None

            # 1. 데이터 로드 및 정리
            if ca_path:
                self.log("CA410 데이터 정리 중...")
                df_ca = self.clean_df(self.read_df(ca_path))
                self.log(f"  CA410 유효 패널: {len(df_ca)}건")

            if lmk_path:
                self.log("LMK6 데이터 정리 중...")
                df_lmk = self.clean_df(self.read_df(lmk_path))
                self.log(f"  LMK6 유효 패널: {len(df_lmk)}건")

            # 2. 개별 정리 파일 생성
            out_path = None
            if make_indiv:
                if df_ca is not None:
                    self.log("CA410 개별 정리 파일(2개 시트) 생성 중...")
                    res = self.save_individual_file(df_ca, "CA410", save_dir)
                    if not both:
                        out_path = res
                if df_lmk is not None:
                    self.log("LMK6 개별 정리 파일(2개 시트) 생성 중...")
                    res = self.save_individual_file(df_lmk, "LMK6", save_dir)
                    if not both:
                        out_path = res

            # 3. 통합 양식 기입 (CA410 + LMK6 둘 다 있을 때만)
            if df_ca is not None and df_lmk is not None:
                if not os.path.exists(temp_path):
                    raise FileNotFoundError(f"양식 파일을 찾을 수 없습니다: {temp_path}")
                if "Panel_ID" not in df_ca.columns or "Panel_ID" not in df_lmk.columns:
                    raise ValueError("CA410/LMK6 양쪽 모두 Panel_ID 컬럼이 필요합니다.")

                self.log("데이터 병합 중...")
                lmk_subset = df_lmk[[c for c in LMK_MERGE_COLS if c in df_lmk.columns]]
                # 겹치는 컬럼이 있으면 pandas 가 _x/_y 접미사를 붙여 양식 기입이 실패한다.
                # 균일도는 LMK6 값이 기준이므로 CA410 쪽 동명 컬럼을 먼저 제거한다.
                overlap = [c for c in lmk_subset.columns if c != "Panel_ID" and c in df_ca.columns]
                if overlap:
                    self.log(f"  CA410 중복 컬럼 제외: {', '.join(overlap)}")
                    df_ca = df_ca.drop(columns=overlap)

                merged = pd.merge(df_ca, lmk_subset, on="Panel_ID", how="outer")
                merged = merged.sort_values("Panel_ID", kind="stable")

                self.log(f"양식 기입 중... ({len(merged)}행)")
                out_path = self.fill_template(merged, temp_path, save_dir)
                self.log(f"✅ 통합 완료: {os.path.basename(out_path)}")

            if out_path:
                self.last_saved_file = out_path
                self.last_saved_folder = save_dir
                self.root.after(0, self._set_result_buttons, True)
                self.root.after(0, self._show_info, "성공", "데이터 정리가 완료되었습니다.")
            else:
                self.log("⚠ 작업은 완료되었으나 생성된 파일이 없습니다 (옵션 확인).")

        except Exception as exc:
            detail = f"{type(exc).__name__}: {exc}"
            self.log(f"❌ 에러: {detail}")
            self.root.after(0, self._show_error, "오류", f"처리 중 에러 발생:\n{detail}")
        finally:
            self.root.after(0, self._set_run_button, "normal", "▶ One-Click 정리 시작")

    def _open_with_os(self, path):
        """OS 기본 프로그램으로 파일/폴더를 연다. (Windows 외 환경도 안전하게)"""
        try:
            opener = getattr(os, "startfile", None)
            if opener is not None:
                opener(path)
                return
            import subprocess
            cmd = "open" if sys.platform == "darwin" else "xdg-open"
            subprocess.Popen([cmd, path])
        except OSError as exc:
            self.log(f"⚠ 열기 실패: {exc}")

    def open_last_file(self):
        if self.last_saved_file:
            self._open_with_os(self.last_saved_file)

    def open_last_folder(self):
        if self.last_saved_folder:
            self._open_with_os(self.last_saved_folder)


def main():
    root = tk.Tk()
    OneClickApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
