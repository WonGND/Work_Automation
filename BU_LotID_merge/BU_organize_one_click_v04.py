import csv
import io
import re
import shutil
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import PatternFill
from PIL import Image

# 처리 대상 이미지 확장자
ALLOWED_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp")
# 파일명에서 LotID/종류(BU, WU)를 뽑기 위한 패턴
LOT_PATTERN = re.compile(r"^(?P<lotid>.+)_(?P<kind>BU|WU)_\d+$", re.IGNORECASE)
DATA_FILE_PATTERN = "LMK6DataLog.csv"
BU_SPEC_MIN = 50.0
WU_SPEC_MIN = 80.0
BU_GRID_COLS = 48
BU_GRID_ROWS = 27
MODEL_NAME_CANDIDATES = (
    "Model_Name", "ModelName", "Model", "RecipeName", "Recipe", "Product", "Product_Name",
)

class PipelineCancelled(Exception):
    pass

def get_resized_xl_image(image_path: Path, max_width_px: int) -> XLImage | None:
    if not image_path.exists():
        return None
    try:
        with Image.open(image_path) as img:
            w, h = img.size
            if w > max_width_px:
                ratio = max_width_px / w
                new_w, new_h = int(w * ratio), int(h * ratio)
                img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
            
            img_byte_arr = io.BytesIO()
            img.save(img_byte_arr, format="PNG")
            img_byte_arr.seek(0)
            return XLImage(img_byte_arr)
    except Exception as e:
        print(f"Error resizing image {image_path}: {e}")
        return None

def print_progress(label: str, current: int, total: int, done: bool = False) -> None:
    if total <= 0: return
    percent = (current / total) * 100
    print(f"{label}: {current}/{total} ({percent:5.1f}%)", flush=True)

def ensure_not_cancelled(cancel_check=None) -> None:
    if cancel_check and cancel_check():
        raise PipelineCancelled("사용자 요청으로 작업이 중지되었습니다.")

def parse_lot_kind(filename: str) -> tuple[str, str]:
    match = LOT_PATTERN.search(filename)
    if match:
        return match.group("lotid"), match.group("kind").upper()
    return filename, "UNKNOWN"

def collect_latest_lotid_folders(integrated_root: Path, cancel_check=None) -> tuple[dict[str, Path], list[dict]]:
    lotid_folders = [d for d in integrated_root.iterdir() if d.is_dir()]
    total = len(lotid_folders)
    print(f"\n[1/7] LotID 폴더 스캔 시작 (대상 폴더: {total}개)")

    latest_by_lotid: dict[str, Path] = {}
    all_rows = []

    for idx, folder in enumerate(lotid_folders, start=1):
        ensure_not_cancelled(cancel_check)
        lot_id = folder.name
        created_time = folder.stat().st_ctime
        
        row = {
            "lot_id": lot_id,
            "folder_path": str(folder),
            "created_time": datetime.fromtimestamp(created_time).strftime("%Y-%m-%d %H:%M:%S"),
        }

        if lot_id not in latest_by_lotid or created_time > latest_by_lotid[lot_id].stat().st_ctime:
            latest_by_lotid[lot_id] = folder

        all_rows.append(row)
        if idx == 1 or idx % 50 == 0 or idx == total:
            print_progress("  스캔 진행", idx, total, done=(idx == total))

    return latest_by_lotid, all_rows

def copy_latest_folders(latest_by_lotid: dict[str, Path], dst_root: Path, cancel_check=None) -> None:
    if dst_root.exists():
        shutil.rmtree(dst_root)
    dst_root.mkdir(parents=True, exist_ok=True)

    total = len(latest_by_lotid)
    print(f"\n[2/7] 최신 LotID 폴더 복사 시작 (대상: {total}개)")

    for idx, (lot_id, src_path) in enumerate(latest_by_lotid.items(), start=1):
        ensure_not_cancelled(cancel_check)
        shutil.copytree(src_path, dst_root / lot_id)
        if idx == 1 or idx % 20 == 0 or idx == total:
            print_progress("  복사 진행", idx, total, done=(idx == total))

def collect_latest_measurements(data_root: Path, cancel_check=None) -> tuple[dict[str, dict], list[dict]]:
    csv_files = list(data_root.rglob(DATA_FILE_PATTERN))
    total = len(csv_files)
    print(f"\n[4/7] 측정 CSV 스캔 시작 (대상 파일: {total}개)")

    latest_measurements: dict[str, dict] = {}
    all_measurement_rows = []

    for idx, csv_path in enumerate(csv_files, start=1):
        ensure_not_cancelled(cancel_check)
        try:
            try:
                df = pd.read_csv(csv_path, encoding="utf-8")
            except:
                df = pd.read_csv(csv_path, encoding="euc-kr")

            model_name = "Unknown"
            for cand in MODEL_NAME_CANDIDATES:
                if cand in df.columns:
                    model_name = str(df[cand].iloc[0]) if not df.empty else "Unknown"
                    break

            for _, row_data in df.iterrows():
                lot_id = str(row_data.get("LotID", ""))
                if not lot_id or lot_id.lower() == "nan": continue
                
                time_str = str(row_data.get("Time", ""))
                try:
                    time_val = datetime.strptime(time_str, "%Y-%m-%d %H:%M:%S")
                except:
                    time_val = datetime.fromtimestamp(csv_path.stat().st_mtime)

                m_row = {
                    "lot_id": lot_id,
                    "judge": str(row_data.get("Judge", "")),
                    "black_uniformity": float(row_data.get("Black_Uniformity", 0)),
                    "white_uniformity": float(row_data.get("White_Uniformity", 0)),
                    "time_raw": time_val,
                    "time_str": time_str,
                    "model_name": model_name,
                }

                if lot_id not in latest_measurements or time_val > latest_measurements[lot_id]["time_raw"]:
                    latest_measurements[lot_id] = m_row
                all_measurement_rows.append(m_row)
        except Exception: pass

        if idx == 1 or idx % 50 == 0 or idx == total:
            print_progress("  CSV 스캔 진행", idx, total, done=(idx == total))

    return latest_measurements, all_measurement_rows

def find_non_black_bbox(img: Image.Image, threshold: int = 12) -> tuple[int, int, int, int] | None:
    arr = np.array(img.convert("L"))
    non_black = np.where(arr > threshold)
    if non_black[0].size == 0 or non_black[1].size == 0: return None
    return (int(np.min(non_black[1])), int(np.min(non_black[0])), int(np.max(non_black[1])), int(np.max(non_black[0])))

def crop_images(merged_root: Path, cropped_root: Path, threshold: int, padding: int, cancel_check=None) -> list[dict]:
    if cropped_root.exists(): shutil.rmtree(cropped_root)
    cropped_root.mkdir(parents=True, exist_ok=True)

    image_files = []
    for ext in ALLOWED_EXTENSIONS:
        image_files.extend(list(merged_root.rglob(f"*{ext}")))
    
    total = len(image_files)
    print(f"\n[5/7] 이미지 크롭 시작 (대상: {total}개, 멀티스레드 활성화)")

    records = []
    lock = threading.Lock()
    processed_count = 0

    def process_one(src: Path):
        nonlocal processed_count
        if cancel_check and cancel_check(): return
        
        rel_path = src.relative_to(merged_root)
        dst = cropped_root / rel_path
        dst.parent.mkdir(parents=True, exist_ok=True)
        lot_id, kind = parse_lot_kind(src.stem)

        try:
            with Image.open(src) as img:
                bbox = find_non_black_bbox(img, threshold=threshold)
                if bbox is None:
                    status = "NO_OBJECT_DETECTED"
                    used_bbox = (0, 0, img.width, img.height)
                    cropped = img.copy()
                else:
                    status = "OK"
                    x_min, y_min, x_max, y_max = bbox
                    used_bbox = (max(0, x_min-padding), max(0, y_min-padding), min(img.width, x_max+padding), min(img.height, y_max+padding))
                    cropped = img.crop(used_bbox)

                if dst.suffix.lower() in (".jpg", ".jpeg") and cropped.mode not in ("RGB", "L"):
                    cropped = cropped.convert("RGB")
                cropped.save(dst)
            
            with lock:
                records.append({"lot_id": lot_id, "kind": kind, "src": src, "dst": dst, "bbox": used_bbox, "status": status})
        except Exception as exc:
            with lock:
                records.append({"lot_id": lot_id, "kind": kind, "src": src, "dst": None, "bbox": None, "status": f"ERROR: {exc}"})
        
        with lock:
            processed_count += 1
            if processed_count == 1 or processed_count % 20 == 0 or processed_count == total:
                print_progress("  크롭 진행", processed_count, total, done=(processed_count == total))

    with ThreadPoolExecutor(max_workers=4) as executor:
        executor.map(process_one, image_files)

    return records

def write_excel(records, excel_path: Path, latest_measurements=None, image_width_px: int = 240, cancel_check=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "결과"
    ws.append(["LotID", "판정", "BU data 수치화", "BU Image", "WU data", "WU Image"])

    grouped = {}
    for rec in records:
        lot_id, kind = rec["lot_id"], rec["kind"]
        grouped.setdefault(lot_id, {"BU": None, "WU": None})
        if kind in ("BU", "WU") and rec["dst"] and not grouped[lot_id][kind]:
            grouped[lot_id][kind] = rec["dst"]

    row_idx = 2
    for lot_id in sorted(grouped.keys()):
        ensure_not_cancelled(cancel_check)
        ws.cell(row=row_idx, column=1, value=lot_id)
        m = (latest_measurements or {}).get(lot_id, {})
        ws.cell(row=row_idx, column=2, value=m.get("judge", ""))
        ws.cell(row=row_idx, column=3, value=m.get("black_uniformity", ""))
        ws.cell(row=row_idx, column=5, value=m.get("white_uniformity", ""))

        max_h = 0
        for col, kind in [("D", "BU"), ("F", "WU")]:
            img_path = grouped[lot_id][kind]
            if img_path and Path(img_path).exists():
                xl_img = get_resized_xl_image(Path(img_path), image_width_px)
                if xl_img:
                    ws.add_image(xl_img, f"{col}{row_idx}")
                    max_h = max(max_h, xl_img.height)

        ws.row_dimensions[row_idx].height = max(25, int(max_h * 0.75))
        row_idx += 1

    # 조건부 서식
    last_row = max(2, row_idx - 1)
    green_fill = PatternFill("solid", fgColor="D1FAE5")
    red_fill = PatternFill("solid", fgColor="FEE2E2")
    yellow_fill = PatternFill("solid", fgColor="FEF3C7")
    
    ws.conditional_formatting.add(f"B2:B{last_row}", FormulaRule(formula=['B2="OK"'], fill=green_fill))
    ws.conditional_formatting.add(f"B2:B{last_row}", FormulaRule(formula=['B2="NG"'], fill=red_fill))
    ws.conditional_formatting.add(f"C2:C{last_row}", CellIsRule(operator="lessThan", formula=[str(BU_SPEC_MIN)], fill=yellow_fill))
    ws.conditional_formatting.add(f"E2:E{last_row}", CellIsRule(operator="lessThan", formula=[str(WU_SPEC_MIN)], fill=yellow_fill))

    # 컬럼 너비
    for col, width in zip("ABCDEF", [26, 12, 14, 36, 14, 36]):
        ws.column_dimensions[col].width = width

    # Visualization 시트
    if latest_measurements:
        v_ws = wb.create_sheet("Visualization")
        v_ws.append(["LotID", "ModelName", "BU", "WU", "Judge", "Time"])
        for lid in sorted(latest_measurements.keys()):
            m = latest_measurements[lid]
            v_ws.append([lid, m.get("model_name", ""), m.get("black_uniformity", 0), m.get("white_uniformity", 0), m.get("judge", ""), m.get("time_str", "")])
        
        if len(latest_measurements) > 1:
            chart = LineChart()
            chart.title = "Uniformity Trend"
            chart.y_axis.title = "Value (%)"
            chart.add_data(Reference(v_ws, min_col=3, min_row=1, max_col=4, max_row=len(latest_measurements)+1), titles_from_data=True)
            chart.set_categories(Reference(v_ws, min_col=1, min_row=2, max_row=len(latest_measurements)+1))
            v_ws.add_chart(chart, "H2")

    print("\n메인 엑셀 저장")
    wb.save(excel_path)

def write_bu_analysis_excel(records, output_path, cancel_check=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "BU_Grid_Analysis"
    ws.append(["LotID", "Min", "Max", "Avg", "Std"])
    
    bu_recs = [r for r in records if r["kind"] == "BU" and r["dst"]]
    for rec in bu_recs:
        ensure_not_cancelled(cancel_check)
        try:
            img = cv2.imread(str(rec["dst"]), cv2.IMREAD_GRAYSCALE)
            if img is not None:
                ws.append([rec["lot_id"], np.min(img), np.max(img), np.mean(img), np.std(img)])
        except: pass
    wb.save(output_path)
    return len(bu_recs)

def run_pipeline(integrated_root: Path, data_root: Path, threshold: int, padding: int, cancel_check=None) -> dict:
    cropped_root = integrated_root.parent / f"{integrated_root.name}_LotID_latest_v1_cropped_v1"
    excel_path = cropped_root / "crop_report.xlsx"
    bu_excel_path = cropped_root / "bu_grid_analysis.xlsx"

    ensure_not_cancelled(cancel_check)
    latest_folders, _ = collect_latest_lotid_folders(integrated_root, cancel_check)
    if not latest_folders: return {}

    merged_root = integrated_root.parent / f"{integrated_root.name}_LotID_latest_v1"
    copy_latest_folders(latest_folders, merged_root, cancel_check)
    
    latest_m, _ = collect_latest_measurements(data_root, cancel_check)
    crop_records = crop_images(merged_root, cropped_root, threshold, padding, cancel_check)
    
    write_excel(crop_records, excel_path, latest_m, cancel_check=cancel_check)
    bu_count = write_bu_analysis_excel(crop_records, bu_excel_path, cancel_check)

    print("\n--- 최종 결과 (v0.4 Optimized) ---")
    print(f"완료! (성공: {sum(1 for r in crop_records if r['status']=='OK')})")

    return {"excel_path": excel_path, "bu_analysis_excel_path": bu_excel_path}

if __name__ == "__main__":
    print("\n--- BU Organize One Click v0.4 Optimized ---")
    ir = Path(input("1) 이미지 통합 폴더: ").strip())
    dr = Path(input("2) 측정 데이터 폴더: ").strip())
    th = int(input("3) 임계값 [12]: ") or 12)
    pd = int(input("4) 패딩 [8]: ") or 8)
    run_pipeline(ir, dr, th, pd)
