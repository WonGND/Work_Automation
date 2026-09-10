import csv
import os
import shutil
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

from bu_common import (
    ALLOWED_EXTENSIONS,
    find_non_black_bbox,
    get_resized_xl_image,
    iter_csv_dict_rows,
    normalize_lot_id,
    parse_lot_kind,
    print_file_created,
    print_progress,
    print_stage,
)

DATA_FILE_PATTERN = "LMK6DataLog.csv"
BU_SPEC_MIN = 50.0
WU_SPEC_MIN = 80.0
MODEL_NAME_CANDIDATES = (
    "Model_Name", "ModelName", "Model", "RecipeName", "Recipe", "Product", "Product_Name",
)
LOT_ID_CANDIDATES = ("Panel_ID", "LotID", "Lot_ID", "PanelID", "Panel Id")

class PipelineCancelled(Exception):
    pass

def ensure_not_cancelled(cancel_check=None) -> None:
    if cancel_check and cancel_check():
        raise PipelineCancelled("사용자 요청으로 작업이 중지되었습니다.")

def first_row_value(row: dict, candidates: tuple[str, ...]) -> str:
    for key in candidates:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""

def parse_lmk_time(value, fallback_path: Path) -> datetime:
    raw = "" if value is None else str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y.%m.%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return datetime.fromtimestamp(fallback_path.stat().st_mtime)

def collect_latest_lotid_folders(
    integrated_root: Path,
    cancel_check=None,
    recursive: bool = True,
    total_steps: int = 5,
) -> tuple[dict[str, Path], list[dict]]:
    """이미지가 직접 든 LotID 폴더를 찾아 이름별 최신 폴더를 고른다.

    recursive=False는 예전 v0.5의 직계 하위 폴더 범위가 필요한 경우를 위한 선택지다.
    """
    candidates = integrated_root.rglob("*") if recursive else integrated_root.iterdir()
    lotid_folders = [
        folder
        for folder in candidates
        if folder.is_dir()
        and any(
            child.is_file() and child.suffix.lower() in ALLOWED_EXTENSIONS
            for child in folder.iterdir()
        )
    ]
    total = len(lotid_folders)
    scope = "전체 하위" if recursive else "직계 하위"
    print_stage(1, total_steps, "LotID 폴더 스캔", f"{scope}, 대상 폴더 {total}개")

    latest_by_lotid: dict[str, Path] = {}
    latest_times: dict[str, tuple[float, float]] = {}
    all_rows: list[dict] = []

    for idx, folder in enumerate(lotid_folders, start=1):
        ensure_not_cancelled(cancel_check)
        lot_id = folder.name
        stat = folder.stat()
        folder_time = (stat.st_ctime, stat.st_mtime)
        
        row = {
            "lot_id": lot_id,
            "folder_path": str(folder),
            "created_time": datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S"),
            "modified_time": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        }

        if lot_id not in latest_by_lotid or folder_time > latest_times[lot_id]:
            latest_by_lotid[lot_id] = folder
            latest_times[lot_id] = folder_time

        all_rows.append(row)
        if idx == 1 or idx % 50 == 0 or idx == total:
            print_progress("  스캔 진행", idx, total, done=(idx == total))

    selected = {str(path) for path in latest_by_lotid.values()}
    for row in all_rows:
        row["selected_latest_final"] = "TRUE" if row["folder_path"] in selected else "FALSE"
    return latest_by_lotid, all_rows

def copy_latest_folders(latest_by_lotid: dict[str, Path], dst_root: Path, cancel_check=None, total_steps: int = 5) -> None:
    dst_root.mkdir(parents=True, exist_ok=True)
    selected_names = set(latest_by_lotid)
    for stale in dst_root.iterdir():
        if stale.name not in selected_names:
            if stale.is_dir():
                shutil.rmtree(stale)
            else:
                stale.unlink()

    items = sorted(latest_by_lotid.items())
    total = len(items)
    print_stage(2, total_steps, "최신 LotID 폴더 병렬 복사", f"대상 {total}개")
    lock = threading.Lock()
    copied = 0

    def copy_one(item: tuple[str, Path]) -> None:
        nonlocal copied
        ensure_not_cancelled(cancel_check)
        lot_id, src_path = item
        destination = dst_root / lot_id
        if destination.exists():
            shutil.rmtree(destination)
        shutil.copytree(src_path, destination)
        with lock:
            copied += 1
            if copied == 1 or copied % 20 == 0 or copied == total:
                print_progress("  복사 진행", copied, total, done=copied == total)

    max_workers = min(16, max(1, len(items)), (os.cpu_count() or 4) + 4)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        list(executor.map(copy_one, items))

def collect_latest_measurements(data_root: Path, cancel_check=None, total_steps: int = 5) -> tuple[dict[str, dict], list[dict]]:
    csv_files = sorted(data_root.rglob(DATA_FILE_PATTERN))
    total = len(csv_files)
    print_stage(3, total_steps, "측정 CSV 스캔", f"대상 파일 {total}개")

    latest_measurements: dict[str, dict] = {}
    all_measurement_rows = []
    if total == 0:
        print("  측정 CSV를 찾지 못해서 판정/BU/WU data는 공란으로 둡니다.")
        return latest_measurements, all_measurement_rows

    for idx, csv_path in enumerate(csv_files, start=1):
        ensure_not_cancelled(cancel_check)
        try:
            for row_data in iter_csv_dict_rows(csv_path):
                lot_id = first_row_value(row_data, LOT_ID_CANDIDATES)
                if not lot_id:
                    continue

                time_str = str(row_data.get("Time", "")).strip()
                time_val = parse_lmk_time(time_str, csv_path)

                m_row = {
                    "lot_id": lot_id,
                    "judge": str(row_data.get("Judge", "")).strip(),
                    "black_uniformity": str(row_data.get("Black_Uniformity", "")).strip(),
                    "white_uniformity": str(row_data.get("White_Uniformity", "")).strip(),
                    "time_raw": time_val,
                    "time_str": time_str,
                    "model_name": first_row_value(row_data, MODEL_NAME_CANDIDATES),
                    "source_file": str(csv_path),
                }

                normalized_lot_id = normalize_lot_id(lot_id)
                if (
                    normalized_lot_id not in latest_measurements
                    or time_val > latest_measurements[normalized_lot_id]["time_raw"]
                ):
                    latest_measurements[normalized_lot_id] = m_row
                all_measurement_rows.append({**m_row, "status": "OK"})
        except (OSError, UnicodeError, csv.Error, ValueError) as exc:
            print(f"  CSV 읽기 실패: {csv_path} ({exc})")
            all_measurement_rows.append(
                {
                    "lot_id": "",
                    "judge": "",
                    "black_uniformity": "",
                    "white_uniformity": "",
                    "time_raw": "",
                    "time_str": "",
                    "model_name": "",
                    "source_file": str(csv_path),
                    "status": f"ERROR: {exc}",
                }
            )

        if idx == 1 or idx % 50 == 0 or idx == total:
            print_progress("  CSV 스캔 진행", idx, total, done=(idx == total))

    return latest_measurements, all_measurement_rows

def find_measurement(lot_id: str, latest_measurements: dict[str, dict]) -> dict:
    normalized = normalize_lot_id(lot_id)
    if normalized in latest_measurements:
        return latest_measurements[normalized]
    candidates = [
        (key, measurement)
        for key, measurement in latest_measurements.items()
        if normalized and (normalized in key or key in normalized)
    ]
    if candidates:
        longest = max(len(key) for key, _ in candidates)
        best = [(key, measurement) for key, measurement in candidates if len(key) == longest]
        if len(best) == 1:
            key, measurement = best[0]
            print(f"  측정값 유사 LotID 연결: {lot_id} -> {key}")
            return measurement
        print(f"  측정값 유사 LotID가 모호하여 연결하지 않음: {lot_id}")
    return {}

def excel_measurement_value(value):
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    try:
        return float(text)
    except ValueError:
        return text

def crop_images(merged_root: Path, cropped_root: Path, threshold: int, padding: int, cancel_check=None, total_steps: int = 5) -> list[dict]:
    if cropped_root.exists():
        shutil.rmtree(cropped_root)
    cropped_root.mkdir(parents=True, exist_ok=True)

    image_files = sorted(
        path
        for path in merged_root.rglob("*")
        if path.is_file() and path.suffix.lower() in ALLOWED_EXTENSIONS
    )
    
    total = len(image_files)
    print_stage(4, total_steps, "이미지 크롭", f"대상 {total}개, 멀티스레드")

    records = []
    lock = threading.Lock()
    processed_count = 0

    def process_one(src: Path) -> None:
        nonlocal processed_count
        ensure_not_cancelled(cancel_check)
        
        rel_path = src.relative_to(merged_root)
        dst = cropped_root / rel_path
        dst.parent.mkdir(parents=True, exist_ok=True)
        lot_id, kind = parse_lot_kind(src.stem)

        try:
            with Image.open(src) as img:
                bbox = find_non_black_bbox(img, threshold=threshold)
                if bbox is None:
                    status = "NO_OBJECT_DETECTED"
                    used_bbox = (0, 0, img.width, img.height)
                    cropped = img.copy()
                else:
                    status = "OK"
                    x_min, y_min, x_max, y_max = bbox
                    used_bbox = (max(0, x_min-padding), max(0, y_min-padding), min(img.width, x_max+padding), min(img.height, y_max+padding))
                    cropped = img.crop(used_bbox)

                if dst.suffix.lower() in (".jpg", ".jpeg") and cropped.mode not in ("RGB", "L"):
                    cropped = cropped.convert("RGB")
                cropped.save(dst)
            
            record = {"lot_id": lot_id, "kind": kind, "src": src, "dst": dst, "bbox": used_bbox, "status": status}
        except Exception as exc:
            record = {"lot_id": lot_id, "kind": kind, "src": src, "dst": None, "bbox": None, "status": f"ERROR: {exc}"}
            print(f"  이미지 크롭 실패: {src} ({exc})")

        with lock:
            records.append(record)
            processed_count += 1
            if processed_count == 1 or processed_count % 20 == 0 or processed_count == total:
                print_progress("  크롭 진행", processed_count, total, done=(processed_count == total))

    max_workers = min(32, (os.cpu_count() or 4) + 4)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        list(executor.map(process_one, image_files))

    records.sort(key=lambda record: str(record["src"]))
    return records

def write_excel(
    records,
    excel_path: Path,
    latest_measurements=None,
    merge_rows=None,
    measurement_rows=None,
    image_width_px: int = 240,
    cancel_check=None,
    total_steps: int = 5,
    excel_step: int = 5,
    analyses=None,
):
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("엑셀 결과 시트를 만들지 못했습니다.")
    ws.title = "결과"
    analysis_by_lot = {item.lot_id: item for item in (analyses or []) if item.status == "OK"}
    header = ["LotID", "판정", "BU data 수치화", "BU Image", "WU data", "WU Image"]
    if analysis_by_lot:
        header += ["Weak 비율(%)", "Weak 집중 영역"]
    ws.append(header)
    print_stage(excel_step, total_steps, "메인 엑셀 작성", f"이미지 기록 {len(records)}개")

    grouped = {}
    for rec in records:
        lot_id, kind = rec["lot_id"], rec["kind"]
        grouped.setdefault(lot_id, {"BU": None, "WU": None})
        if kind in ("BU", "WU") and rec["dst"] and not grouped[lot_id][kind]:
            grouped[lot_id][kind] = rec["dst"]

    row_idx = 2
    for lot_id in sorted(grouped.keys()):
        ensure_not_cancelled(cancel_check)
        ws.cell(row=row_idx, column=1, value=lot_id)
        m = find_measurement(lot_id, latest_measurements or {})
        ws.cell(row=row_idx, column=2, value=m.get("judge", ""))
        ws.cell(row=row_idx, column=3, value=excel_measurement_value(m.get("black_uniformity", "")))
        ws.cell(row=row_idx, column=5, value=excel_measurement_value(m.get("white_uniformity", "")))

        max_h = 0
        for col, kind in [("D", "BU"), ("F", "WU")]:
            img_path = grouped[lot_id][kind]
            if img_path and Path(img_path).exists():
                xl_img = get_resized_xl_image(Path(img_path), image_width_px)
                if xl_img:
                    ws.add_image(xl_img, f"{col}{row_idx}")
                    max_h = max(max_h, xl_img.height)

        analysis = analysis_by_lot.get(lot_id)
        if analysis is not None:
            ws.cell(row=row_idx, column=7, value=round(analysis.weak_ratio * 100, 4))
            ws.cell(row=row_idx, column=8, value=analysis.dominant_zone)

        ws.row_dimensions[row_idx].height = max(25, int(max_h * 0.75))
        row_idx += 1

    # 조건부 서식
    last_row = max(2, row_idx - 1)
    green_fill = PatternFill("solid", fgColor="D1FAE5")
    red_fill = PatternFill("solid", fgColor="FEE2E2")
    yellow_fill = PatternFill("solid", fgColor="FEF3C7")
    
    ws.conditional_formatting.add(f"B2:B{last_row}", FormulaRule(formula=['B2="OK"'], fill=green_fill))
    ws.conditional_formatting.add(f"B2:B{last_row}", FormulaRule(formula=['B2="NG"'], fill=red_fill))
    ws.conditional_formatting.add(f"C2:C{last_row}", CellIsRule(operator="lessThan", formula=[str(BU_SPEC_MIN)], fill=yellow_fill))
    ws.conditional_formatting.add(f"E2:E{last_row}", CellIsRule(operator="lessThan", formula=[str(WU_SPEC_MIN)], fill=yellow_fill))

    # 컬럼 너비
    for col, width in zip("ABCDEFGH", [26, 12, 14, 36, 14, 36, 14, 16]):
        ws.column_dimensions[col].width = width

    detail_ws = wb.create_sheet("처리_상세")
    detail_ws.append(["구분", "LotID", "종류", "상태", "원본 경로", "결과 경로", "기타"])
    for record in records:
        detail_ws.append(
            [
                "CROP",
                record["lot_id"],
                record["kind"],
                record["status"],
                str(record["src"]),
                "" if record["dst"] is None else str(record["dst"]),
                "" if record["bbox"] is None else str(record["bbox"]),
            ]
        )
    for row in merge_rows or []:
        detail_ws.append(
            [
                "MERGE",
                row["lot_id"],
                "",
                "SELECTED" if row.get("selected_latest_final") == "TRUE" else "EXCLUDED",
                row["folder_path"],
                "",
                f"created={row['created_time']}, modified={row['modified_time']}",
            ]
        )
    for row in measurement_rows or []:
        detail_ws.append(
            [
                "MEASURE",
                row.get("lot_id", ""),
                "",
                row.get("status", "OK"),
                row.get("source_file", ""),
                "",
                f"time={row.get('time_str', '')}, BU={row.get('black_uniformity', '')}, WU={row.get('white_uniformity', '')}",
            ]
        )
    for column, width in zip("ABCDEFG", [12, 26, 10, 28, 60, 60, 48]):
        detail_ws.column_dimensions[column].width = width

    write_analysis_sheets(wb, analyses or [])

    print("  엑셀 파일 저장 중...")
    wb.save(excel_path)
    print_file_created(excel_path, "메인 엑셀")

def write_merge_report(rows: list[dict], output_root: Path) -> Path:
    report_path = output_root / "merge_report.csv"
    fieldnames = ["lot_id", "folder_path", "created_time", "modified_time", "selected_latest_final"]
    with report_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print_file_created(report_path, "병합 리포트")
    return report_path

COLOR_MAP_MAX_PRODUCTS = 200


def write_color_map_sheet(wb, analyses) -> None:
    """제품별 BU 이미지를 격자로 줄여 셀 배경색으로 다시 그린다.

    이미지를 그대로 넣으면 눈으로만 볼 수 있지만, 셀로 그리면 엑셀에서 확대해
    특정 칸을 짚어보거나 옆에 수치를 붙여 비교할 수 있다. weak 로 잡힌 칸은
    굵은 테두리로 감싸 색만으로 구분되지 않는 경계를 드러낸다.
    """
    from bu_image_analysis import GRID_COLS, GRID_ROWS

    drawable = [item for item in analyses if item.status == "OK" and item.color_grid]
    if not drawable:
        return

    sheet = wb.create_sheet("색분포_맵")
    truncated = drawable[:COLOR_MAP_MAX_PRODUCTS]
    edge = Side(style="medium", color="111827")
    weak_border = Border(left=edge, right=edge, top=edge, bottom=edge)

    sheet.cell(row=1, column=1, value="굵은 테두리 = weak point 로 판정된 칸").font = Font(bold=True)
    anchor = 3
    for analysis in truncated:
        title = sheet.cell(row=anchor, column=1)
        title.value = (
            f"{analysis.lot_id}  ·  weak {analysis.weak_ratio * 100:.2f}%"
            f"  ·  집중 영역 {analysis.dominant_zone}"
        )
        title.font = Font(bold=True)

        grid_top = anchor + 1
        for column in range(GRID_COLS):
            header = sheet.cell(row=grid_top, column=column + 2, value=column + 1)
            header.font = Font(size=7)
            header.alignment = Alignment(horizontal="center")

        for row in range(GRID_ROWS):
            label = sheet.cell(row=grid_top + 1 + row, column=1, value=row + 1)
            label.font = Font(size=7)
            label.alignment = Alignment(horizontal="center")
            sheet.row_dimensions[grid_top + 1 + row].height = 12

            for column in range(GRID_COLS):
                cell = sheet.cell(row=grid_top + 1 + row, column=column + 2)
                color = analysis.color_grid[row][column]
                if color is not None:
                    cell.fill = PatternFill("solid", fgColor=color)
                if analysis.weak_grid[row][column]:
                    cell.border = weak_border

        anchor = grid_top + GRID_ROWS + 3

    sheet.column_dimensions["A"].width = 4
    for column in range(GRID_COLS):
        sheet.column_dimensions[get_column_letter(column + 2)].width = 2.4

    if len(drawable) > COLOR_MAP_MAX_PRODUCTS:
        sheet.cell(
            row=anchor,
            column=1,
            value=(
                f"제품이 많아 앞 {COLOR_MAP_MAX_PRODUCTS}개만 그렸습니다. "
                f"(전체 {len(drawable)}개)"
            ),
        ).font = Font(bold=True, color="B45309")


def analyze_bu_records(
    crop_records: list[dict],
    cancel_check=None,
    step: int = 5,
    total_steps: int = 6,
) -> list:
    """크롭 결과에서 BU 이미지를 골라 색 분포와 weak point 를 계산한다."""
    from bu_image_analysis import analyze_bu_image

    targets = [
        record
        for record in crop_records
        if record["kind"] == "BU" and record["dst"] and Path(record["dst"]).exists()
    ]
    total = len(targets)
    print_stage(step, total_steps, "BU Image 분석", f"대상 {total}개")
    if total == 0:
        print("  분석할 BU 이미지가 없어 건너뜁니다.")
        return []

    analyses = []
    for index, record in enumerate(targets, start=1):
        ensure_not_cancelled(cancel_check)
        analysis = analyze_bu_image(Path(record["dst"]), record["lot_id"])
        analyses.append(analysis)
        if analysis.status != "OK":
            print(f"  분석 실패: {record['lot_id']} ({analysis.status})")
        if index == 1 or index % 10 == 0 or index == total:
            print_progress("  분석 진행", index, total, done=(index == total))

    ok_count = sum(1 for item in analyses if item.status == "OK")
    weak_count = sum(
        1 for item in analyses if item.status == "OK" and item.dominant_zone != "없음"
    )
    print(f"  분석 완료 (성공 {ok_count}개, weak 검출 {weak_count}개)")
    return analyses


def write_analysis_sheets(wb, analyses: list) -> None:
    """분석 결과 시트를 이미 열려 있는 통합 워크북에 붙인다.

    분석이 전부 실패했다면 통계 시트는 빈 껍데기가 되므로, 실패 사유만 남는
    요약 시트 한 장으로 끝낸다.
    """
    from bu_image_analysis import (
        COLOR_NAMES,
        WEAK_SEVERITY_MIN,
        ZONE_NAMES,
        aggregate_zone_stats,
    )

    if not analyses:
        return

    summary_ws = wb.create_sheet("제품별_요약")
    summary_ws.append(
        ["LotID", "상태", "평균 심각도", "Weak 비율(%)", "Weak 집중 영역", "집중 영역 비율(%)"]
        + [f"{name}(%)" for name in COLOR_NAMES]
    )
    for analysis in analyses:
        summary_ws.append(
            [
                analysis.lot_id,
                analysis.status,
                round(analysis.mean_severity, 5),
                round(analysis.weak_ratio * 100, 4),
                analysis.dominant_zone,
                round(analysis.dominant_zone_ratio * 100, 4),
            ]
            + [round(analysis.color_ratios.get(name, 0.0) * 100, 4) for name in COLOR_NAMES]
        )
    for column, width in zip("ABCDEFGHIJKL", [40, 10, 14, 14, 16, 18, 12, 12, 12, 12, 12, 12]):
        summary_ws.column_dimensions[column].width = width
    summary_ws.freeze_panes = "A2"

    if not any(item.status == "OK" for item in analyses):
        return

    zone_ws = wb.create_sheet("영역별_Weak")
    zone_ws.append(["LotID"] + list(ZONE_NAMES))
    for analysis in analyses:
        zone_ws.append(
            [analysis.lot_id]
            + [round(analysis.zone_weak_ratios.get(name, 0.0) * 100, 4) for name in ZONE_NAMES]
        )
    zone_ws.column_dimensions["A"].width = 40
    for column in "BCDEFGHIJ":
        zone_ws.column_dimensions[column].width = 10
    zone_ws.freeze_panes = "B2"

    last_zone_row = max(2, zone_ws.max_row)
    zone_ws.conditional_formatting.add(
        f"B2:J{last_zone_row}",
        CellIsRule(
            operator="greaterThan",
            formula=["0"],
            fill=PatternFill("solid", fgColor="FEE2E2"),
        ),
    )

    stats = aggregate_zone_stats(analyses)
    total_ws = wb.create_sheet("영역별_누적")
    total_ws.append(["영역", "평균 Weak 비율(%)", "최대 Weak 비율(%)", "검출 제품 수", "검출 비율(%)"])
    for name in ZONE_NAMES:
        item = stats[name]
        total_ws.append(
            [
                name,
                round(item["평균 weak 비율"] * 100, 4),
                round(item["최대 weak 비율"] * 100, 4),
                item["검출 제품 수"],
                round(item["검출 비율"] * 100, 2),
            ]
        )
    for column, width in zip("ABCDE", [10, 20, 20, 14, 14]):
        total_ws.column_dimensions[column].width = width

    criteria_ws = wb.create_sheet("판정_기준")
    criteria_ws.append(["항목", "값", "설명"])
    for row in (
        ["흰색", "1.00", "가장 낮은 데이터. 계측기 글자로 판정되면 제외"],
        ["노랑", "0.80", "두 번째로 나쁨"],
        ["주황", "0.60", "세 번째로 나쁨"],
        ["빨강", "0.40", "네 번째로 나쁨"],
        ["청록", "0.15", "초록보다 낮음"],
        ["초록", "0.00", "가장 양호"],
        ["Weak 기준", f"{WEAK_SEVERITY_MIN:.2f} 이상", "주황 이상을 weak point 로 판정"],
    ):
        criteria_ws.append(row)
    for column, width in zip("ABC", [14, 16, 52]):
        criteria_ws.column_dimensions[column].width = width

    write_color_map_sheet(wb, analyses)

def run_pipeline(
    integrated_root: Path,
    data_root: Path,
    threshold: int,
    padding: int,
    cancel_check=None,
    analyze_bu_images: bool = False,
) -> dict:
    cropped_root = integrated_root.parent / f"{integrated_root.name}_LotID_latest_v1_cropped_v1"
    excel_path = cropped_root / "crop_report.xlsx"
    total_steps = 6 if analyze_bu_images else 5

    ensure_not_cancelled(cancel_check)
    latest_folders, merge_rows = collect_latest_lotid_folders(
        integrated_root, cancel_check, total_steps=total_steps
    )
    if not latest_folders:
        raise RuntimeError("LotID 폴더를 찾지 못했습니다. 이미지 통합 폴더 구조를 확인하세요.")

    merged_root = integrated_root.parent / f"{integrated_root.name}_LotID_latest_v1"
    copy_latest_folders(latest_folders, merged_root, cancel_check, total_steps=total_steps)
    merge_report_path = write_merge_report(merge_rows, merged_root)

    latest_m, measurement_rows = collect_latest_measurements(
        data_root, cancel_check, total_steps=total_steps
    )
    crop_records = crop_images(
        merged_root, cropped_root, threshold, padding, cancel_check, total_steps=total_steps
    )

    analyses = []
    if analyze_bu_images:
        analyses = analyze_bu_records(
            crop_records,
            cancel_check=cancel_check,
            step=5,
            total_steps=total_steps,
        )

    write_excel(
        crop_records,
        excel_path,
        latest_m,
        merge_rows,
        measurement_rows,
        cancel_check=cancel_check,
        total_steps=total_steps,
        excel_step=total_steps,
        analyses=analyses,
    )

    ok_analyses = [item for item in analyses if item.status == "OK"]
    weak_products = sum(1 for item in ok_analyses if item.dominant_zone != "없음")

    success_count = sum(1 for record in crop_records if record["status"] == "OK")
    error_count = sum(1 for record in crop_records if record["status"].startswith("ERROR"))
    print("\n--- 최종 결과 ---")
    print(f"완료! (크롭 성공: {success_count}, 오류: {error_count})")
    if analyze_bu_images:
        print(f"BU Image 분석: {len(ok_analyses)}개 분석, weak 검출 {weak_products}개")

    return {
        "merged_root": merged_root,
        "cropped_root": cropped_root,
        "excel_path": excel_path,
        "merge_report_path": merge_report_path,
        "latest_lotids": len(latest_folders),
        "crop_records": len(crop_records),
        "crop_ok": success_count,
        "crop_error": error_count,
        "analysis_excel_path": excel_path if analyze_bu_images else None,
        "analyzed_images": len(ok_analyses),
        "weak_products": weak_products,
    }

if __name__ == "__main__":
    print("\n--- BU Organize One Click ---")
    ir = Path(input("1) 이미지 통합 폴더: ").strip())
    dr = Path(input("2) 측정 데이터 폴더: ").strip())
    th = int(input("3) 임계값 [12]: ") or 12)
    padding = int(input("4) 패딩 [20]: ") or 20)
    analyze = (input("5) BU Image 분석 실행 (y/N): ").strip().lower() == "y")
    run_pipeline(ir, dr, th, padding, analyze_bu_images=analyze)
